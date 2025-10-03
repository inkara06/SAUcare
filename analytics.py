import pandas as pd
import matplotlib.pyplot as plt
import plotly.express as px
import plotly.graph_objects as go
from sqlalchemy import create_engine
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import ColorScaleRule
import os
from datetime import datetime

# Database Configuration
DB_CONFIG = {
    'user': 'inkar',
    'password': '0000',
    'host': 'localhost',
    'port': '5432',
    'database': 'hospital management'
}

def get_db_connection():
    """Create database connection"""
    try:
        connection_string = f"postgresql://{DB_CONFIG['user']}:{DB_CONFIG['password']}@{DB_CONFIG['host']}:{DB_CONFIG['port']}/{DB_CONFIG['database']}"
        engine = create_engine(connection_string)
        print("✅ Successfully connected to database")
        return engine
    except Exception as e:
        print(f"❌ Error connecting to database: {e}")
        return None

def create_directories():
    """Create necessary directories for charts and exports"""
    os.makedirs('charts', exist_ok=True)
    os.makedirs('exports', exist_ok=True)
    print("📁 Created directories: charts/ and exports/")

# ============ VISUALIZATION 1: PIE CHART ============
def create_pie_chart(engine):
    """Distribution of appointments by status"""
    query = """
    SELECT 
        a.status,
        COUNT(*) as count
    FROM appointments a
    JOIN patients p ON a.patient_id = p.patient_id
    GROUP BY a.status
    ORDER BY count DESC;
    """
    
    df = pd.read_sql(query, engine)
    print(f"\n📊 PIE CHART - Appointment Status Distribution")
    print(f"Rows returned: {len(df)}")
    print(df)
    
    plt.figure(figsize=(10, 7))
    colors = ['#66b3ff', '#99ff99', '#ffcc99', '#ff9999']
    plt.pie(df['count'], labels=df['status'], autopct='%1.1f%%', 
            colors=colors, startangle=90)
    plt.title('Distribution of Appointments by Status', fontsize=16, fontweight='bold')
    plt.axis('equal')
    plt.savefig('charts/01_pie_appointment_status.png', dpi=300, bbox_inches='tight')
    plt.close()
    print("✅ Saved: charts/01_pie_appointment_status.png")
    
    return df

# ============ VISUALIZATION 2: BAR CHART ============
def create_bar_chart(engine):
    """Appointment volume by doctor specialization"""
    query = """
    SELECT 
        d.specialization,
        COUNT(a.appointment_id) as appointment_count,
        COUNT(DISTINCT a.patient_id) as unique_patients
    FROM appointments a
    JOIN doctors d ON a.doctor_id = d.doctor_id
    JOIN patients p ON a.patient_id = p.patient_id
    GROUP BY d.specialization
    ORDER BY appointment_count DESC;
    """
    
    df = pd.read_sql(query, engine)
    print(f"\n📊 BAR CHART - Appointments by Doctor Specialization")
    print(f"Rows returned: {len(df)}")
    print(df.head(10))
    
    plt.figure(figsize=(12, 6))
    plt.bar(df['specialization'], df['appointment_count'], color='steelblue', edgecolor='black')
    plt.xlabel('Specialization', fontsize=12, fontweight='bold')
    plt.ylabel('Number of Appointments', fontsize=12, fontweight='bold')
    plt.title('Appointment Volume by Doctor Specialization', fontsize=14, fontweight='bold')
    plt.xticks(rotation=45, ha='right')
    plt.grid(axis='y', alpha=0.3)
    plt.tight_layout()
    plt.savefig('charts/02_bar_specialization.png', dpi=300, bbox_inches='tight')
    plt.close()
    print("✅ Saved: charts/02_bar_specialization.png")
    
    return df

# ============ VISUALIZATION 3: HORIZONTAL BAR CHART ============
def create_horizontal_bar_chart(engine):
    """Average treatment cost by treatment type"""
    query = """
    SELECT 
        t.treatment_type,
        COUNT(t.treatment_id) as treatment_count,
        AVG(t.cost) as avg_cost,
        SUM(t.cost) as total_revenue
    FROM treatments t
    JOIN appointments a ON t.appointment_id = a.appointment_id
    JOIN patients p ON a.patient_id = p.patient_id
    GROUP BY t.treatment_type
    HAVING COUNT(t.treatment_id) > 0
    ORDER BY avg_cost DESC
    LIMIT 10;
    """
    
    df = pd.read_sql(query, engine)
    print(f"\n📊 HORIZONTAL BAR CHART - Average Cost by Treatment Type")
    print(f"Rows returned: {len(df)}")
    print(df)
    
    plt.figure(figsize=(12, 8))
    plt.barh(df['treatment_type'], df['avg_cost'], color='coral', edgecolor='black')
    plt.xlabel('Average Cost ($)', fontsize=12, fontweight='bold')
    plt.ylabel('Treatment Type', fontsize=12, fontweight='bold')
    plt.title('Average Treatment Cost by Type', fontsize=14, fontweight='bold')
    plt.grid(axis='x', alpha=0.3)
    plt.tight_layout()
    plt.savefig('charts/03_horizontal_bar_treatment_cost.png', dpi=300, bbox_inches='tight')
    plt.close()
    print("✅ Saved: charts/03_horizontal_bar_treatment_cost.png")
    
    return df

# ============ VISUALIZATION 4: LINE CHART ============
def create_line_chart(engine):
    """Monthly appointment trends"""
    query = """
    SELECT 
        TO_CHAR(a.appointment_date, 'YYYY-MM') as month,
        COUNT(a.appointment_id) as appointment_count,
        COUNT(DISTINCT a.patient_id) as unique_patients,
        COUNT(DISTINCT a.doctor_id) as active_doctors
    FROM appointments a
    JOIN patients p ON a.patient_id = p.patient_id
    JOIN doctors d ON a.doctor_id = d.doctor_id
    GROUP BY TO_CHAR(a.appointment_date, 'YYYY-MM')
    ORDER BY month;
    """
    
    df = pd.read_sql(query, engine)
    print(f"\n📊 LINE CHART - Monthly Appointment Trends")
    print(f"Rows returned: {len(df)}")
    print(df)
    
    plt.figure(figsize=(14, 6))
    plt.plot(df['month'], df['appointment_count'], marker='o', linewidth=2, 
             label='Total Appointments', color='blue')
    plt.plot(df['month'], df['unique_patients'], marker='s', linewidth=2, 
             label='Unique Patients', color='green')
    plt.xlabel('Month', fontsize=12, fontweight='bold')
    plt.ylabel('Count', fontsize=12, fontweight='bold')
    plt.title('Monthly Appointment and Patient Trends', fontsize=14, fontweight='bold')
    plt.legend()
    plt.xticks(rotation=45, ha='right')
    plt.grid(True, alpha=0.3)
    plt.tight_layout()
    plt.savefig('charts/04_line_monthly_trends.png', dpi=300, bbox_inches='tight')
    plt.close()
    print("✅ Saved: charts/04_line_monthly_trends.png")
    
    return df

# ============ VISUALIZATION 5: HISTOGRAM ============
def create_histogram(engine):
    """Distribution of treatment costs"""
    query = """
    SELECT 
        t.cost,
        t.treatment_type,
        a.appointment_date
    FROM treatments t
    JOIN appointments a ON t.appointment_id = a.appointment_id
    JOIN billing b ON t.treatment_id = b.treatment_id
    WHERE t.cost IS NOT NULL AND t.cost > 0;
    """
    
    df = pd.read_sql(query, engine)
    print(f"\n📊 HISTOGRAM - Treatment Cost Distribution")
    print(f"Rows returned: {len(df)}")
    print(f"Cost range: ${df['cost'].min():.2f} - ${df['cost'].max():.2f}")
    print(f"Average cost: ${df['cost'].mean():.2f}")
    
    plt.figure(figsize=(12, 6))
    plt.hist(df['cost'], bins=20, color='skyblue', edgecolor='black', alpha=0.7)
    plt.xlabel('Treatment Cost ($)', fontsize=12, fontweight='bold')
    plt.ylabel('Frequency', fontsize=12, fontweight='bold')
    plt.title('Distribution of Treatment Costs', fontsize=14, fontweight='bold')
    plt.axvline(df['cost'].mean(), color='red', linestyle='--', linewidth=2, label=f'Mean: ${df["cost"].mean():.2f}')
    plt.legend()
    plt.grid(axis='y', alpha=0.3)
    plt.tight_layout()
    plt.savefig('charts/05_histogram_cost_distribution.png', dpi=300, bbox_inches='tight')
    plt.close()
    print("✅ Saved: charts/05_histogram_cost_distribution.png")
    
    return df

# ============ VISUALIZATION 6: SCATTER PLOT ============
def create_scatter_plot(engine):
    """Doctor experience vs patient load"""
    query = """
    SELECT 
        d.doctor_id,
        d.first_name || ' ' || d.last_name as doctor_name,
        d.specialization,
        d.years_experience,
        COUNT(a.appointment_id) as total_appointments,
        COUNT(DISTINCT a.patient_id) as unique_patients
    FROM doctors d
    LEFT JOIN appointments a ON d.doctor_id = a.doctor_id
    LEFT JOIN patients p ON a.patient_id = p.patient_id
    GROUP BY d.doctor_id, d.first_name, d.last_name, d.specialization, d.years_experience
    HAVING COUNT(a.appointment_id) > 0
    ORDER BY d.years_experience;
    """
    
    df = pd.read_sql(query, engine)
    print(f"\n📊 SCATTER PLOT - Doctor Experience vs Patient Load")
    print(f"Rows returned: {len(df)}")
    print(df.head(10))
    
    plt.figure(figsize=(12, 7))
    scatter = plt.scatter(df['years_experience'], df['total_appointments'], 
                          s=df['unique_patients']*10, alpha=0.6, c=df['unique_patients'],
                          cmap='viridis', edgecolors='black')
    plt.xlabel('Years of Experience', fontsize=12, fontweight='bold')
    plt.ylabel('Total Appointments', fontsize=12, fontweight='bold')
    plt.title('Doctor Experience vs Patient Load\n(Bubble size = Unique Patients)', 
              fontsize=14, fontweight='bold')
    plt.colorbar(scatter, label='Unique Patients')
    plt.grid(True, alpha=0.3)
    plt.tight_layout()
    plt.savefig('charts/06_scatter_experience_load.png', dpi=300, bbox_inches='tight')
    plt.close()
    print("✅ Saved: charts/06_scatter_experience_load.png")
    
    return df

# ============ INTERACTIVE TIME SLIDER (PLOTLY) ============
def create_plotly_time_slider(engine):
    """Interactive chart with time slider - Monthly appointment trends by specialization"""
    query = """
    SELECT 
        TO_CHAR(a.appointment_date, 'YYYY-MM') as month,
        d.specialization,
        COUNT(a.appointment_id) as appointment_count,
        COUNT(DISTINCT a.patient_id) as unique_patients
    FROM appointments a
    JOIN doctors d ON a.doctor_id = d.doctor_id
    JOIN patients p ON a.patient_id = p.patient_id
    GROUP BY TO_CHAR(a.appointment_date, 'YYYY-MM'), d.specialization
    ORDER BY month, specialization;
    """
    
    df = pd.read_sql(query, engine)
    print(f"\n🎬 INTERACTIVE TIME SLIDER - Monthly Trends by Specialization")
    print(f"Rows returned: {len(df)}")
    
    fig = px.bar(df, 
                 x='specialization', 
                 y='appointment_count',
                 animation_frame='month',
                 color='specialization',
                 title='Monthly Appointments by Doctor Specialization (Interactive Slider)',
                 labels={'appointment_count': 'Number of Appointments', 'specialization': 'Specialization'},
                 height=600)
    
    fig.update_layout(showlegend=False)
    fig.show()
    print("✅ Interactive chart displayed in browser")
    
    return df

# ============ EXPORT TO EXCEL WITH FORMATTING ============
def export_to_excel(dataframes_dict, filename):
    """Export multiple dataframes to Excel with professional formatting"""
    filepath = f'exports/{filename}'
    
    # Write dataframes to Excel
    with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
        total_rows = 0
        for sheet_name, df in dataframes_dict.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            total_rows += len(df)
    
    # Apply formatting
    wb = load_workbook(filepath)
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        
        # Freeze panes (header row)
        ws.freeze_panes = "A2"
        
        # Add auto filter
        ws.auto_filter.ref = ws.dimensions
        
        # Apply conditional formatting to numeric columns
        for col_idx, col in enumerate(ws.iter_cols(min_row=2, max_row=ws.max_row), start=1):
            col_letter = col[0].column_letter
            
            # Check if column is numeric
            try:
                if col[0].value and isinstance(col[0].value, (int, float)):
                    # Apply color scale rule
                    rule = ColorScaleRule(
                        start_type="min", start_color="FFAA0000",
                        mid_type="percentile", mid_value=50, mid_color="FFFFFF00",
                        end_type="max", end_color="FF00AA00"
                    )
                    ws.conditional_formatting.add(f"{col_letter}2:{col_letter}{ws.max_row}", rule)
            except:
                pass
        
        # Style header row
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        for cell in ws[1]:
            cell.fill = header_fill
    
    wb.save(filepath)
    
    num_sheets = len(dataframes_dict)
    print(f"\n📊 Created file {filename}, {num_sheets} sheets, {total_rows} rows")
    return filepath

def main():
    """Main execution function"""
    print("="*70)
    print("🏥 HOSPITAL MANAGEMENT ANALYTICS - ASSIGNMENT 2")
    print("="*70)
    
    # Create directories
    create_directories()
    
    # Connect to database
    engine = get_db_connection()
    if not engine:
        print("\n⚠️  Cannot connect to database. Please check your connection settings.")
        return
    
    # Store dataframes for Excel export
    export_data = {}
    
    try:
        # Create all visualizations
        print("\n" + "="*70)
        print("CREATING VISUALIZATIONS")
        print("="*70)
        
        export_data['Appointment_Status'] = create_pie_chart(engine)
        export_data['Specialization'] = create_bar_chart(engine)
        export_data['Treatment_Costs'] = create_horizontal_bar_chart(engine)
        export_data['Monthly_Trends'] = create_line_chart(engine)
        export_data['Cost_Distribution'] = create_histogram(engine)
        export_data['Doctor_Analysis'] = create_scatter_plot(engine)
        
        # Create interactive plot
        print("\n" + "="*70)
        print("CREATING INTERACTIVE VISUALIZATION")
        print("="*70)
        create_plotly_time_slider(engine)
        
        # Export to Excel
        print("\n" + "="*70)
        print("EXPORTING TO EXCEL")
        print("="*70)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        export_to_excel(export_data, f'hospital_analytics_{timestamp}.xlsx')
        
        print("\n" + "="*70)
        print("✅ ALL TASKS COMPLETED SUCCESSFULLY!")
        print("="*70)
        print("\n📁 Check the following directories:")
        print("   - charts/     (6 visualization images)")
        print("   - exports/    (Excel file with formatted data)")
        
    except Exception as e:
        print(f"\n❌ Error during execution: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    main()
